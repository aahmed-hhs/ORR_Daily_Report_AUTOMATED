import datetime as dt
from pathlib import Path

from docx import Document
from PIL import Image

from powerbi_report.document import fill_template


def test_fill_template(tmp_path: Path):
    tpl = tmp_path / "tpl.docx"
    d = Document()
    d.add_paragraph("Report for {{DATE}}")
    d.add_paragraph("{{DASHBOARD}}")
    d.save(tpl)

    img = tmp_path / "img.png"
    Image.new("RGB", (200, 100), "white").save(img)

    out = fill_template(tpl, img, tmp_path / "out.docx", run_date=dt.date(2026, 1, 15))
    doc = Document(out)
    texts = [p.text for p in doc.paragraphs]
    assert "Report for January 15, 2026" in texts
    assert "{{DASHBOARD}}" not in texts
    assert len(doc.inline_shapes) == 1
