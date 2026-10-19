from __future__ import annotations

import datetime as dt
from pathlib import Path

from docx import Document
from docx.shared import Inches
from docx.text.paragraph import Paragraph


def _iter_paragraphs(doc):
    """Yield every paragraph in the body and inside tables."""
    yield from doc.paragraphs
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs


def _replace_text(para: Paragraph, old: str, new: str) -> bool:
    hit = False
    for run in para.runs:
        if old in run.text:
            run.text = run.text.replace(old, new)
            hit = True
    if not hit and old in para.text:
        # placeholder split across runs: collapse into the first run
        full = para.text.replace(old, new)
        for run in para.runs:
            run.text = ""
        para.runs[0].text = full
        hit = True
    return hit


def _place_image(para: Paragraph, image: Path, width_in: float) -> None:
    for run in para.runs:
        run.text = ""
    para.add_run().add_picture(str(image), width=Inches(width_in))


def fill_template(
    template: Path,
    image: Path,
    out_path: Path,
    *,
    image_placeholder: str = "{{DASHBOARD}}",
    date_placeholder: str = "{{DATE}}",
    date_format: str = "%B %d, %Y",
    image_width_inches: float = 6.5,
    run_date: dt.date | None = None,
) -> Path:
    run_date = run_date or dt.date.today()
    doc = Document(str(template))
    placed = False
    for para in _iter_paragraphs(doc):
        if image_placeholder in para.text:
            _place_image(para, image, image_width_inches)
            placed = True
        if date_placeholder in para.text:
            _replace_text(para, date_placeholder, run_date.strftime(date_format))
    if not placed:
        raise SystemExit(f"Could not find {image_placeholder} in {template}")
    out_path.parent.mkdir(parents=True, exist_ok=True)
    try:
        doc.save(str(out_path))
    except PermissionError:
        # Usually the file is open in Word. Save next to it instead of failing.
        alt = out_path
        for n in range(1, 100):
            alt = out_path.with_name(f"{out_path.stem}_{n}{out_path.suffix}")
            if not alt.exists():
                break
        print(f"{out_path.name} is locked (open in Word?); saving as {alt.name} instead")
        doc.save(str(alt))
        out_path = alt
    print(f"Wrote {out_path}")
    return out_path
