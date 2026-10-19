"""Rewrite a Power BI 'Export data' workbook into the classic layout:
headers in row 1, data, blank row, then the 'Applied filters' note at the bottom;
sheet named 'Export'; date columns shown as long dates."""
from __future__ import annotations

import datetime as dt
from pathlib import Path

import openpyxl

LONG_DATE = "[$-F800]dddd\\,\\ mmmm\\ dd\\,\\ yyyy"


def to_classic_layout(path: Path, sheet_name: str = "Export", date_format: str = LONG_DATE) -> Path:
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    rows = [list(r) for r in ws.iter_rows(values_only=True)]

    # Power BI (new layout) puts an 'Applied filters:' block at the top followed by a blank row.
    note = None
    if rows and isinstance(rows[0][0], str) and rows[0][0].lower().startswith("applied filters"):
        note = rows[0][0]
        rows = rows[1:]
        while rows and all(v in (None, "") for v in rows[0]):
            rows = rows[1:]
    # Already classic? (note at the bottom) - leave the body alone but still normalise.
    elif rows and isinstance(rows[-1][0], str) and rows[-1][0].lower().startswith("applied filters"):
        note = rows[-1][0]
        rows = rows[:-1]
        while rows and all(v in (None, "") for v in rows[-1]):
            rows = rows[:-1]

    if not rows:
        return path
    ncols = max(len(r) for r in rows)
    header, body = rows[0], rows[1:]

    out = openpyxl.Workbook()
    o = out.active
    o.title = sheet_name
    o.append(["" if v is None else v for v in header] + [""] * (ncols - len(header)))
    date_cols: set[int] = set()
    for r in body:
        r = ["" if v is None else v for v in r] + [""] * (ncols - len(r))
        o.append(r)
        for i, v in enumerate(r, start=1):
            if isinstance(v, (dt.datetime, dt.date)):
                date_cols.add(i)
    if note:
        o.append([""] * ncols)
        o.append([note] + [""] * (ncols - 1))
    for c in date_cols:
        for row in o.iter_rows(min_row=2, max_row=1 + len(body), min_col=c, max_col=c):
            for cell in row:
                if isinstance(cell.value, (dt.datetime, dt.date)):
                    cell.number_format = date_format
    out.save(path)
    return path